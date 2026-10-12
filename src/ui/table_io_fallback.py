# -*- coding: utf-8 -*-
"""表格双通道读写（R183/F2）：openpyxl 在则走 .xlsx，不在则自动降级 stdlib csv。

背景（深查③ T9/P7，实机复核）：`pyproject.toml` 声明了 `openpyxl`，但真机 venv
里没装 —— 两处「导出/导入 Excel」（TopView 布局表、Patch 配接表）都在
`_ensure_openpyxl` 处直接弹错误框返回 = **功能隐形不可用**（用户以为能用，
点了只有报错）。任务要求：不得新增依赖，导出统一走 stdlib csv 降级（utf-8-sig
带 BOM，Excel 双击直接开），并在状态栏提示"走的哪条路"。

设计：
  * 「有没有 openpyxl」收敛成一个探测点（`openpyxl_pack`，带缓存）；
  * `write_table` / `read_table` 是唯一出入口：返回 `(path, fmt)` / `(headers, rows, fmt)`，
    `fmt ∈ {"xlsx", "csv"}`，调用方据此提示用户；
  * 列结构、sheet 名、说明页文案由调用方给出 → 与旧实现逐行同构，只换承载格式；
  * 提示走 `notify_user`：状态栏 showMessage 为主、toast 为辅，两者都取不到时静默
    （无 Qt 的纯逻辑测试路径不依赖提示成功）。

无 Qt 依赖歧义：本模块仅 `notify_user` 触碰 Qt，且整段 try/except 包裹。
"""
from __future__ import annotations

import csv
import io
import os
from typing import Iterable, List, Optional, Sequence, Tuple

XLSX_SUFFIX = ".xlsx"
CSV_SUFFIX = ".csv"
#: csv 编码：utf-8-sig（BOM）—— Excel/WPS 双击直接识别中文，不出现乱码。
CSV_ENCODING = "utf-8-sig"
#: 读取时的编码回退链（与 `import_patch_csv.read_patch_csv` 同款顺序）。
READ_ENCODINGS = ("utf-8-sig", "gb18030", "cp936", "utf-8")

_PACK_CACHE: Optional[object] = None
_PACK_TRIED = False


def openpyxl_pack():
    """返回 (Workbook, load_workbook, Font, PatternFill, Alignment, Comment) 或 None。

    探测一次后缓存（含失败结果）：缺库是常态，不必每次点按钮都试 import。
    """
    global _PACK_CACHE, _PACK_TRIED
    if _PACK_TRIED:
        return _PACK_CACHE
    _PACK_TRIED = True
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.comments import Comment
        _PACK_CACHE = (Workbook, load_workbook, Font, PatternFill, Alignment, Comment)
    except Exception:
        _PACK_CACHE = None
    return _PACK_CACHE


def openpyxl_available() -> bool:
    return openpyxl_pack() is not None


def reset_probe_cache() -> None:
    """测试用：清掉探测缓存（允许 monkeypatch 后重新探测）。"""
    global _PACK_CACHE, _PACK_TRIED
    _PACK_CACHE = None
    _PACK_TRIED = False


def csv_path_for(path: str) -> str:
    """把 .xlsx 目标路径改写为同名 .csv（已是 .csv 则原样）。"""
    root, ext = os.path.splitext(str(path))
    if ext.lower() == CSV_SUFFIX:
        return str(path)
    return root + CSV_SUFFIX


def write_csv(path: str, headers: Sequence, rows: Iterable[Sequence]) -> str:
    """写 utf-8-sig CSV（带 BOM，Excel 直接开）。返回最终路径。"""
    out = csv_path_for(path)
    with io.open(out, "w", encoding=CSV_ENCODING, newline="") as f:
        w = csv.writer(f)
        w.writerow(list(headers))
        for row in rows:
            w.writerow(list(row))
    return out


def _decode_text(path: str) -> str:
    """按编码回退链读全文；全失败退 gb18030 容错（不抛，交给上层报格式错误）。"""
    for enc in READ_ENCODINGS:
        try:
            with io.open(path, "r", encoding=enc, newline="") as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with io.open(path, "r", encoding="gb18030", errors="replace", newline="") as f:
        return f.read()


def read_csv(path: str) -> Tuple[List[str], List[List[str]]]:
    """读 CSV，返回 (表头, 数据行)。编码回退链见 READ_ENCODINGS。

    `#` 开头的行是我们自己写的「说明」注释行（xlsx 路径有独立工作表），
    导入时直接丢弃 —— 数据行不会以 `#` 开头，判据安全。
    """
    text = _decode_text(path)
    reader = csv.reader(io.StringIO(text))
    raw = [list(r) for r in reader
           if not (r and str(r[0]).strip().startswith("#"))]
    if not raw:
        return [], []
    headers = [str(c).strip() for c in raw[0]]
    rows = [list(r) for r in raw[1:] if any(str(c).strip() for c in r)]
    return headers, rows


def _write_xlsx(path: str, pack, headers: Sequence, rows: Iterable[Sequence], *,
                sheet_title: str, header_fill: str, widths: Optional[Sequence[int]],
                help_title: str, help_lines: Optional[Sequence[str]],
                header_comment: str) -> str:
    """openpyxl 路径的写盘（表头样式 + 列宽 + 说明工作表，观感与旧实现一致）。"""
    Workbook, _load_workbook, Font, PatternFill, Alignment, Comment = pack
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(headers))
    for cell in ws[1]:
        _style_header_cell(cell, pack, header_fill, header_comment)
    for row in rows:
        ws.append(list(row))
    for idx, width in enumerate(widths or (), 1):
        if idx <= 26:
            ws.column_dimensions[chr(64 + idx)].width = width
    if help_lines:
        help_ws = wb.create_sheet(help_title)
        for line in help_lines:
            help_ws.append([line])
    wb.save(path)
    return str(path)


def _write_csv(path: str, headers: Sequence, rows: Iterable[Sequence],
               help_lines: Optional[Sequence[str]]) -> str:
    """csv 路径的写盘：utf-8-sig（BOM，Excel 双击直接开中文不乱码）。"""
    out = csv_path_for(path)
    with io.open(out, "w", encoding=CSV_ENCODING, newline="") as f:
        w = csv.writer(f)
        w.writerow(list(headers))
        for row in rows:
            w.writerow(list(row))
        for line in (help_lines or ()):
            w.writerow(["# " + str(line)])
    return out


def _style_header_cell(cell, pack, header_fill: str, header_comment: str) -> None:
    """表头样式（白字 + 主题填充色 + 居中 + 批注），与旧实现逐项一致。"""
    _W, _L, Font, PatternFill, Alignment, Comment = pack
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=header_fill)
    cell.alignment = Alignment(horizontal="center")
    if header_comment:
        cell.comment = Comment(header_comment, "DMX")


def write_table(path: str, headers: Sequence, rows: Iterable[Sequence], *,
                sheet_title: str, header_fill: str = "305496",
                widths: Optional[Sequence[int]] = None,
                help_title: str = "说明",
                help_lines: Optional[Sequence[str]] = None,
                header_comment: str = "字段说明见本工作簿「说明」工作表。") -> Tuple[str, str]:
    """写表格：openpyxl 可用且目标为 .xlsx → 真 xlsx；否则降级 csv。

    返回 (最终路径, "xlsx"|"csv")。csv 路径把「说明」写成表尾 `#` 注释行
    （导入端自动忽略）—— 表头仍在第 1 行，Excel 双击打开即是一张干净的台账。
    """
    pack = openpyxl_pack() if str(path).lower().endswith(XLSX_SUFFIX) else None
    if pack is None:
        return _write_csv(path, headers, rows, help_lines), "csv"
    return _write_xlsx(path, pack, headers, rows, sheet_title=sheet_title,
                       header_fill=header_fill, widths=widths,
                       help_title=help_title, help_lines=help_lines,
                       header_comment=header_comment), "xlsx"


def read_table(path: str, *, sheet_title: str = "") -> Tuple[List[str], List[List[str]], str]:
    """读表格：.xlsx（且 openpyxl 可用）→ 指定 sheet/活动 sheet；否则按 csv 读。

    返回 (表头, 数据行, "xlsx"|"csv")。
    """
    is_xlsx = str(path).lower().endswith(XLSX_SUFFIX)
    pack = openpyxl_pack() if is_xlsx else None
    if pack is None:
        headers, rows = read_csv(path)
        return headers, rows, "csv"
    _Workbook, load_workbook, _Font, _PatternFill, _Alignment, _Comment = pack
    wb = load_workbook(path, data_only=True)
    if sheet_title and sheet_title in wb.sheetnames:
        ws = wb[sheet_title]
    else:
        ws = wb.active
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    if not raw:
        return [], [], "xlsx"
    headers = [str(c).strip() if c is not None else "" for c in raw[0]]
    rows = []
    for r in raw[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        rows.append([None if v is None else v for v in r])
    return headers, rows, "xlsx"


def cell_int(row: Sequence, idx: int, default: int = 0) -> int:
    """按列取值转 int（缺列/空值/非法 → default）。"""
    if idx < 0 or idx >= len(row):
        return default
    v = row[idx]
    if v is None or str(v).strip() == "":
        return default
    try:
        return int(float(str(v).strip()))
    except Exception:
        return default


def cell_str(row: Sequence, idx: int, default: str = "") -> str:
    """按列取值转去空白 str（缺列/空值 → default）。"""
    if idx < 0 or idx >= len(row):
        return default
    v = row[idx]
    if v is None:
        return default
    return str(v).strip() or default


def cell_float(row: Sequence, idx: int, default: float = 0.0) -> float:
    """按列取值转 float（缺列/空值/非法 → default）。"""
    if idx < 0 or idx >= len(row):
        return default
    v = row[idx]
    if v is None or str(v).strip() == "":
        return default
    try:
        return float(str(v).strip())
    except Exception:
        return default


def fmt_note(fmt: str, target: str = "导出") -> str:
    """给用户看的一行"走了哪条路"说明。"""
    if fmt == "xlsx":
        return f"已{target}为 Excel 工作簿 (.xlsx)。"
    return ("未检测到 openpyxl，已自动降级为 CSV 表格 (.csv，UTF-8 带 BOM，"
            "可直接用 Excel/WPS 打开)。")


def notify_user(parent, message: str, detail: str = "", kind: str = "info") -> None:
    """状态栏 + toast 提示（任一路径不可用则跳过；绝不因提示失败影响主流程）。"""
    if parent is None:
        return
    try:
        win = parent.window() if hasattr(parent, "window") else None
        bar = win.statusBar() if win is not None and hasattr(win, "statusBar") else None
        if bar is not None:
            text = message if not detail else f"{message}  {detail}"
            bar.showMessage(text, 8000)
    except Exception:
        pass
    try:
        from .toast_notification import show_info_toast, show_success_toast
        fn = show_success_toast if kind == "success" else show_info_toast
        fn(parent, message, detail or None, duration_ms=4000)
    except Exception:
        pass
